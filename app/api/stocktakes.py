import io
from datetime import date
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.crud import stocktake as crud
from app.core.database import get_db
from app.schemas.stocktake import (
    LatestQuantityOut,
    StatusRow,
    StocktakeCreate,
    StocktakeOut,
    StocktakeUpdate,
)

router = APIRouter(prefix="/api/stocktakes", tags=["stocktakes"])

# 2026-10-28: 현황 엑셀 다운로드 컬럼 정의 — 바코드 추적 여부(센서 바코드
# 조인 테이블 데이터)는 제외하고, 품목/현재고 정보만 내려준다.
_STATUS_EXPORT_COLUMNS = [
    ("품목명", 24),
    ("품번", 14),
    ("구분", 12),
    ("단위", 10),
    ("위치", 12),
    ("규격약식", 20),
    ("현재고", 10),
    ("최근 조사수량(수동)", 18),
    ("최근 조사일자", 14),
]


def _size_label(item) -> str:
    """static/js/common.js::itemSizeLabel과 동일한 "규격약식(폭 x 길이 x 높이)" 포맷."""
    abbr = item.size_abbr.name if item.size_abbr else ""
    parts = [item.size_width, item.size_length, item.size_height]
    has_detail = any(p is not None for p in parts)
    detail = " x ".join(str(p) if p is not None else "-" for p in parts) if has_detail else ""
    if abbr and detail:
        return f"{abbr}({detail})"
    return abbr or detail


def _build_status_workbook(rows) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "재고현황"

    header_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill(fill_type="solid", start_color="354649", end_color="354649")
    header_align = Alignment(horizontal="center", vertical="center")
    for col, (label, width) in enumerate(_STATUS_EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    numeric_align = Alignment(horizontal="right")
    for row_idx, (item, quantity, survey_date, _barcode_count, current_stock) in enumerate(rows, start=2):
        values = [
            item.item_name,
            item.item_no or "",
            item.category.name,
            item.unit.name,
            item.location.name if item.location else "",
            _size_label(item),
            current_stock,
            quantity,
            survey_date,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col in (7, 8):
                cell.alignment = numeric_align
            if col == 9 and value is not None:
                cell.number_format = "yyyy-mm-dd"
        ws.row_dimensions[row_idx].height = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@router.get(
    "",
    response_model=list[StocktakeOut],
    summary="재고조사 기록 목록 조회",
    description="날짜 범위(`date_from`/`date_to`)와 품목 검색/구분/위치로 필터링한 재고조사 기록(수량 스냅샷) 목록. 최신순 정렬.",
)
def list_stocktakes(
    search: str | None = None,
    category_id: int | None = None,
    location_id: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    is_active: bool | None = True,
    limit: int = 50,
    offset: int = 0,
    db: Session = Depends(get_db),
):
    return crud.get_stocktakes(
        db,
        search=search,
        category_id=category_id,
        location_id=location_id,
        date_from=date_from,
        date_to=date_to,
        is_active=is_active,
        limit=limit,
        offset=offset,
    )


@router.get(
    "/status",
    response_model=list[StatusRow],
    summary="현황 조회 (품목별 현재고 요약)",
    description=(
        "전용 테이블 없이, 전체 활성 품목 각각에 가장 최근 재고조사 기록과 "
        "활성 센서 바코드 개수를 집계해서 보여준다. 품목별 '진짜 보유 수량'은 "
        "각 행의 current_stock 필드를 봐야 한다 — last_quantity만 보면 바코드 "
        "추적 품목에서는 실제 수량보다 적게 보일 수 있다."
    ),
)
def status_view(
    search: str | None = None,
    category_id: int | None = None,
    location_id: int | None = None,
    limit: int = 50,
    offset: int = 0,
    db: Session = Depends(get_db),
):
    rows = crud.get_status(
        db, search=search, category_id=category_id, location_id=location_id,
        limit=limit, offset=offset,
    )
    return [
        StatusRow(
            item=item,
            last_quantity=quantity,
            last_survey_date=survey_date,
            barcode_count=count,
            current_stock=current_stock,
        )
        for item, quantity, survey_date, count, current_stock in rows
    ]


@router.get(
    "/status/export",
    summary="현황 엑셀 다운로드",
    description=(
        "현황 화면과 동일한 필터(search/category_id/location_id)를 적용한 전체 "
        "결과를 .xlsx 파일로 내려준다. 센서 바코드 집계(바코드 추적 여부)는 "
        "별도 조인 테이블 데이터라 엑셀에는 포함하지 않는다."
    ),
)
def export_status(
    search: str | None = None,
    category_id: int | None = None,
    location_id: int | None = None,
    db: Session = Depends(get_db),
):
    # 2026-10-28: 화면은 페이지네이션(limit)이 있지만 다운로드는 전체 결과가
    # 필요하므로 매우 큰 limit으로 호출 (get_status 자체는 그대로 재사용).
    rows = crud.get_status(
        db, search=search, category_id=category_id, location_id=location_id,
        limit=100000, offset=0,
    )
    buffer = _build_status_workbook(rows)
    filename = f"재고현황_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get(
    "/latest/{item_id}",
    response_model=LatestQuantityOut,
    summary="품목의 가장 최근 조사 수량 조회",
)
def latest_quantity(item_id: int, db: Session = Depends(get_db)):
    """Backs the 재고 등록 form's quantity prefill: once an item is picked,
    show its current known stock rather than starting from blank/zero."""
    return LatestQuantityOut(quantity=crud.get_latest_quantity(db, item_id))


@router.post(
    "",
    response_model=StocktakeOut,
    status_code=201,
    summary="재고조사 기록 등록",
    description="새로운 수량 스냅샷을 기록한다. survey_date를 생략하면 오늘 날짜로 저장된다.",
)
def create_stocktake(payload: StocktakeCreate, db: Session = Depends(get_db)):
    data = payload.model_dump()
    if data["survey_date"] is None:
        data["survey_date"] = date.today()
    return crud.create_stocktake(db, **data)


@router.put(
    "/{record_id}",
    response_model=StocktakeOut,
    summary="재고조사 기록 수정",
    description="요청 본문에 보낸 필드만 수정된다(부분 수정).",
    responses={404: {"description": "해당 id의 재고조사기록이 없음"}},
)
def update_stocktake(record_id: int, payload: StocktakeUpdate, db: Session = Depends(get_db)):
    row = crud.get_stocktake(db, record_id)
    if row is None:
        raise HTTPException(404, "재고조사기록을 찾을 수 없습니다")
    return crud.update_stocktake(db, row, **payload.model_dump(exclude_unset=True))


@router.delete(
    "/{record_id}",
    status_code=204,
    summary="재고조사 기록 삭제 (논리 삭제)",
    responses={404: {"description": "해당 id의 재고조사기록이 없음"}},
)
def delete_stocktake(record_id: int, db: Session = Depends(get_db)):
    row = crud.get_stocktake(db, record_id)
    if row is None:
        raise HTTPException(404, "재고조사기록을 찾을 수 없습니다")
    crud.deactivate_stocktake(db, row)
